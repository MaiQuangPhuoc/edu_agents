"""
Đối chiếu cột 'real' trong sheet 'itemlist' (OMEGA_itemlist_processed2.xlsx)
với tên file ảnh trong thư mục D:\rikai\thumnail
-> in ra các giá trị real KHÔNG có ảnh tương ứng (missing)
"""

import os
import openpyxl

EXCEL_PATH = r"D:\rikai\OMEGA_itemlist_processed2.xlsx"
THUMB_DIR = r"D:\rikai\thumnail"
SHEET_NAME = "itemlist"
COL_REAL = "real"


def find_header_col(ws, header_name):
    for cell in ws[1]:
        if cell.value and str(cell.value).strip() == header_name:
            return cell.column
    raise ValueError(f"Không tìm thấy cột '{header_name}' ở dòng header (dòng 1).")


def get_real_values(excel_path):
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"Không tìm thấy sheet '{SHEET_NAME}'. Các sheet có: {wb.sheetnames}")

    ws = wb[SHEET_NAME]
    col_real = find_header_col(ws, COL_REAL)

    values = []
    skipped_strike = 0
    for row in range(2, ws.max_row + 1):
        cell = ws.cell(row=row, column=col_real)
        v = cell.value

        # Bỏ qua ô có định dạng gạch ngang (strikethrough)
        if cell.font is not None and cell.font.strike:
            skipped_strike += 1
            continue

        if v is not None and str(v).strip():
            values.append((row, str(v).strip()))

    print(f"Đã bỏ qua {skipped_strike} dòng có real bị gạch ngang.")
    return values

def get_extra_thumbnails(real_values, thumb_names):
    """Trả về danh sách đầy đủ các tên ảnh KHÔNG khớp với bất kỳ giá trị real nào."""
    real_set = {r for _row, r in real_values}
    extra = sorted(thumb_names - real_set)
    return extra


def print_extra_thumbnails(real_values, thumb_names):
    extra = get_extra_thumbnails(real_values, thumb_names)
    print(f"Số ảnh KHÔNG khớp real nào: {len(extra)}")
    for name in extra:
        print(f"  - {name}")
    return extra

def get_thumbnail_names(thumb_dir):
    """Trả về dict {stem (không đuôi): tên file đầy đủ (có đuôi)}."""
    names = {}
    for fname in os.listdir(thumb_dir):
        full_path = os.path.join(thumb_dir, fname)
        if os.path.isfile(full_path):
            stem, _ext = os.path.splitext(fname)
            names[stem.strip()] = fname
    return names


def delete_extra_thumbnails(real_values, thumb_names_dict, thumb_dir, dry_run=True):
    """Xóa các ảnh không khớp real nào. dry_run=True chỉ in ra, không xóa thật."""
    real_set = {r for _row, r in real_values}
    extra_stems = sorted(set(thumb_names_dict.keys()) - real_set)

    print(f"Số ảnh sẽ xóa: {len(extra_stems)}")
    for stem in extra_stems:
        fname = thumb_names_dict[stem]
        full_path = os.path.join(thumb_dir, fname)
        if dry_run:
            print(f"  [DRY-RUN] sẽ xóa: {fname}")
        else:
            os.remove(full_path)
            print(f"  Đã xóa: {fname}")

    return extra_stems


def main():
    real_values = get_real_values(EXCEL_PATH)
    thumb_names_dict = get_thumbnail_names(THUMB_DIR)
    thumb_names = set(thumb_names_dict.keys())

    missing_thumb = [(row, real) for row, real in real_values if real not in thumb_names]

    print(f"Tổng số giá trị real (chưa gạch): {len(real_values)}")
    print(f"Tổng số ảnh trong thumbnail: {len(thumb_names)}\n")

    print(f"[1] Số real KHÔNG có ảnh: {len(missing_thumb)}")
    for row, real in missing_thumb:
        print(f"  - Dòng {row}: {real}")

    print()
    # Bước 1: chạy dry-run trước để kiểm tra danh sách sẽ xóa
    delete_extra_thumbnails(real_values, thumb_names_dict, THUMB_DIR, dry_run=True)

    # Bước 2: nếu danh sách đúng, đổi dry_run=False rồi chạy lại để xóa thật
    # delete_extra_thumbnails(real_values, thumb_names_dict, THUMB_DIR, dry_run=False)


if __name__ == "__main__":
    main()

 
 