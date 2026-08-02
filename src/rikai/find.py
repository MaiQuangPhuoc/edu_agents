"""
Xử lý sheet 'itemlist' trong file OMEGA_itemlist.xlsx
Logic:
1. Trích số trong dấu ngoặc () ở cột "Item name & Supporting text"
   -> nếu có -> ghi ĐÈ vào cột "real"
2. Nếu cột 1 KHÔNG có số -> xét cột "Model number":
   - Nếu chỉ có 1 giá trị -> bỏ qua, KHÔNG đổi cột real
   - Nếu có nhiều giá trị (cách nhau bằng xuống dòng \n) -> so khớp với
     giá trị real đang có sẵn:
       + khớp 1 giá trị -> chốt real = giá trị đó (giữ nguyên)
       + không khớp giá trị nào -> KHÔNG đổi, log ra để kiểm tra tay
"""

import re
import sys
import openpyxl

SHEET_NAME = "itemlist"
COL_ITEM = "Item name & Supporting text"
COL_MODEL = "Model number"
COL_REAL = "real"

# Regex lấy số trong ngoặc: hỗ trợ số dạng 232.112.432 hoặc 23.3
PAREN_NUM_RE = re.compile(r"\(([\d]+(?:\.[\d]+)*)\)")


def extract_number_in_parens(text):
    """Lấy số cuối cùng nằm trong ngoặc () của chuỗi text."""
    if not text:
        return None
    matches = PAREN_NUM_RE.findall(str(text))
    return matches[-1].strip() if matches else None


def split_model_values(text):
    """Tách các giá trị trong cột Model number, phân tách bằng xuống dòng."""
    if not text:
        return []
    parts = [p.strip() for p in str(text).split("\n") if p.strip()]
    return parts


def find_header_col(ws, header_name):
    """Tìm chỉ số cột theo tên header ở dòng 1."""
    for cell in ws[1]:
        if cell.value and str(cell.value).strip() == header_name:
            return cell.column
    raise ValueError(f"Không tìm thấy cột '{header_name}' ở dòng header (dòng 1).")


def process(input_path, output_path):
    wb = openpyxl.load_workbook(input_path)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"Không tìm thấy sheet '{SHEET_NAME}'. Các sheet có: {wb.sheetnames}")

    ws = wb[SHEET_NAME]

    col_item = find_header_col(ws, COL_ITEM)
    col_model = find_header_col(ws, COL_MODEL)
    col_real = find_header_col(ws, COL_REAL)

    ambiguous_rows = []   # log các dòng không tự xử lý được
    updated_rows = []

    for row in range(2, ws.max_row + 1):
        item_text = ws.cell(row=row, column=col_item).value
        model_text = ws.cell(row=row, column=col_model).value
        real_cell = ws.cell(row=row, column=col_real)

        num1 = extract_number_in_parens(item_text)

        if num1:
            # Trường hợp 1: cột 1 có số -> ghi đè real
            real_cell.value = num1
            updated_rows.append((row, "cot1", num1))
            continue

        # Trường hợp cột 1 không có số -> xét cột Model number
        model_values = split_model_values(model_text)

        if len(model_values) <= 1:
            # Chỉ 1 giá trị -> bỏ qua, không đổi real
            continue

        # Nhiều giá trị -> so khớp với real hiện có
        current_real = str(real_cell.value).strip() if real_cell.value is not None else None

        if current_real and current_real in model_values:
            real_cell.value = current_real  # chốt (không đổi giá trị, chỉ xác nhận)
            updated_rows.append((row, "chot", current_real))
        else:
            ambiguous_rows.append((row, model_values, current_real))

    wb.save(output_path)

    print(f"Đã xử lý xong. File output: {output_path}")
    print(f"Số dòng được cập nhật/chốt: {len(updated_rows)}")
    if ambiguous_rows:
        print(f"\n⚠ {len(ambiguous_rows)} dòng KHÔNG tự xử lý được (cần kiểm tra tay):")
        for row, values, real in ambiguous_rows:
            print(f"  - Dòng {row}: Model number = {values}, real hiện tại = {real}")


if __name__ == "__main__":
    input_path = r"D:\rikai\OMEGA_itemlist_processed.xlsx"
    output_path = r"D:\rikai\OMEGA_itemlist_processed2.xlsx"
    process(input_path, output_path)