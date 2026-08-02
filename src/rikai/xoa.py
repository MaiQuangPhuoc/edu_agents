"""
Xóa các ảnh trong D:\rikai\thumnail không khớp với bất kỳ giá trị nào
trong cột 'real' (sheet itemlist, bỏ qua ô bị gạch ngang).
"""

import os
import openpyxl

EXCEL_PATH = r"D:\rikai\OMEGA_itemlist_processed2.xlsx"
THUMB_DIR = r"D:\rikai\thumnail"
SHEET_NAME = "itemlist"
COL_REAL = "real"


def find_header_col(ws, header_name):
    for cell in ws[1]:
        if cell.value and str(cell.value).strip() == header_name:
            return cell.column
    raise ValueError(f"Không tìm thấy cột '{header_name}' ở dòng header (dòng 1).")


def get_real_set(excel_path):
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb[SHEET_NAME]
    col_real = find_header_col(ws, COL_REAL)

    real_set = set()
    for row in range(2, ws.max_row + 1):
        cell = ws.cell(row=row, column=col_real)
        if cell.font is not None and cell.font.strike:
            continue
        v = cell.value
        if v is not None and str(v).strip():
            real_set.add(str(v).strip())
    return real_set


def delete_extra_thumbnails(thumb_dir, real_set):
    deleted = []
    for fname in os.listdir(thumb_dir):
        full_path = os.path.join(thumb_dir, fname)
        if not os.path.isfile(full_path):
            continue
        stem, _ext = os.path.splitext(fname)
        if stem.strip() not in real_set:
            os.remove(full_path)
            deleted.append(fname)
            print(f"Đã xóa: {fname}")

    print(f"\nTổng số ảnh đã xóa: {len(deleted)}")
    return deleted


if __name__ == "__main__":
    real_set = get_real_set(EXCEL_PATH)
    delete_extra_thumbnails(THUMB_DIR, real_set)